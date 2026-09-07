"""
FinOps Multicloud
Construido por Wanderlei Grizolli Junior, Sr. Solution Engineer.
Baseado no Microsoft FinOps toolkit (https://github.com/microsoft/finops-toolkit).

Exportacao: relatorio executivo em PDF e planilha completa em Excel.

Os dois sao gerados no servidor, com as mesmas agregacoes da interface e respeitando os
filtros ativos. O PDF traz graficos, texto narrativo e as tabelas que sustentam a
conversa. O Excel traz tudo em abas, para quem quer trabalhar o numero.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone

import threading

import matplotlib

matplotlib.use("Agg")
import matplotlib.ticker as mticker  # noqa: E402
from matplotlib.figure import Figure  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.enums import TA_LEFT  # noqa: E402
from reportlab.lib.pagesizes import A4, landscape  # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import cm  # noqa: E402
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import sobre  # noqa: E402

AZUL = "#0078D4"
AZUL_ESCURO = "#005A9E"
CINZA = "#52658A"
CORES_NUVEM = {"Microsoft Azure": "#0078D4", "Amazon Web Services": "#FF9900", "Google Cloud": "#34A853", "Oracle Cloud": "#C74634"}
PALETA = ["#0078D4", "#00B7C3", "#8B5CF6", "#2FBF71", "#FFB900", "#F7630C", "#50B0F0", "#E5484D", "#C239B3", "#498205"]


def _moeda(v, m="USD"):
    s = {"USD": "US$", "BRL": "R$", "EUR": "EUR"}.get(m, m)
    return f"{s} {v:,.2f}"


def _mes(m):
    if not m or len(m) < 7:
        return m
    nomes = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]
    return f"{nomes[int(m[5:7]) - 1]}/{m[2:4]}"


# ================================================================== graficos (matplotlib)
# API orientada a objeto (Figure direto, sem pyplot): o pyplot guarda estado global e nao e
# seguro quando duas exportacoes rodam ao mesmo tempo em threads diferentes. A trava e uma
# segunda garantia, barata, para a renderizacao.
_TRAVA_GRAFICOS = threading.Lock()
_FMT_MILHAR = mticker.FuncFormatter(lambda v, _: f"{v / 1e6:.1f}M" if v >= 1e6 else f"{v / 1000:.0f}k" if v >= 1000 else f"{v:.0f}")


def _fig(largura=9.5, altura=3.4):
    fig = Figure(figsize=(largura, altura), dpi=150)
    ax = fig.subplots()
    ax.ticklabel_format(style="plain", axis="both", useOffset=False)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_color("#DCE5F0")
    ax.spines["bottom"].set_color("#DCE5F0")
    ax.tick_params(colors=CINZA, labelsize=8)
    ax.yaxis.grid(True, color="#E8EEF7", linestyle="--", linewidth=.7)
    ax.set_axisbelow(True)
    return fig, ax


def _png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png", bbox_inches="tight", facecolor="white")
    fig.clear()
    buf.seek(0)
    return buf


def graf_mensal(mensal, moeda):
    fig, ax = _fig()
    x = [_mes(m["mes"]) for m in mensal]
    ax.fill_between(x, [m["efetivo"] for m in mensal], alpha=.15, color=AZUL)
    ax.plot(x, [m["efetivo"] for m in mensal], color=AZUL, linewidth=2.2, marker="o", markersize=4, label="Custo efetivo")
    ax.plot(x, [m["faturado"] for m in mensal], color="#00B7C3", linewidth=1.6, linestyle="--", label="Faturado")
    ax.legend(frameon=False, fontsize=8, loc="upper left")
    ax.set_title("Evolução mensal", loc="left", fontsize=10, color="#10203A", fontweight="bold")
    ax.yaxis.set_major_formatter(_FMT_MILHAR)
    return _png(fig)


def graf_barras(itens, titulo, cor_por_nome=None, altura=3.4):
    fig, ax = _fig(altura=altura)
    itens = list(reversed(itens[:12]))
    nomes = [i["nome"][:34] for i in itens]
    vals = [i["efetivo"] for i in itens]
    cores = [(cor_por_nome or {}).get(i["nome"], AZUL) for i in itens]
    ax.barh(nomes, vals, color=cores, height=.62)
    for y, v in enumerate(vals):
        ax.text(v, y, f"  {v / 1000:.1f}k" if v >= 1000 else f"  {v:.0f}", va="center", fontsize=7.5, color=CINZA)
    ax.set_title(titulo, loc="left", fontsize=10, color="#10203A", fontweight="bold")
    ax.xaxis.grid(True, color="#E8EEF7", linestyle="--", linewidth=.7)
    ax.yaxis.grid(False)
    ax.xaxis.set_major_formatter(_FMT_MILHAR)
    return _png(fig)


def graf_rosca(itens, titulo, cor_por_nome=None):
    fig = Figure(figsize=(4.6, 3.4), dpi=150)
    ax = fig.subplots()
    itens = itens[:8]
    vals = [i["efetivo"] for i in itens]
    cores = [(cor_por_nome or {}).get(i["nome"], PALETA[k % len(PALETA)]) for k, i in enumerate(itens)]
    ax.pie(vals, colors=cores, startangle=90, wedgeprops={"width": .38, "edgecolor": "white"})
    ax.legend([f"{i['nome'][:26]}  {i['efetivo'] / sum(vals):.0%}" for i in itens], loc="center left", bbox_to_anchor=(1, .5),
              frameon=False, fontsize=7.5)
    ax.set_title(titulo, loc="left", fontsize=10, color="#10203A", fontweight="bold")
    return _png(fig)


def graf_previsao(prev, moeda):
    fig, ax = _fig()
    h = prev.get("historico", [])[-60:]
    p = prev.get("previsao", [])[:90]
    if h:
        ax.plot(range(len(h)), [x["valor"] for x in h], color=AZUL, linewidth=1.8, label="Realizado")
    if p:
        xs = range(len(h), len(h) + len(p))
        ax.plot(xs, [x["valor"] for x in p], color="#8B5CF6", linewidth=1.8, linestyle="--", label="Previsão")
        ax.fill_between(xs, [x["min"] for x in p], [x["max"] for x in p], color="#8B5CF6", alpha=.15, label="Faixa de 80%")
    ax.legend(frameon=False, fontsize=8, loc="upper left")
    ax.set_title("Custo diário: realizado e previsão de 90 dias", loc="left", fontsize=10, color="#10203A", fontweight="bold")
    ax.set_xticks([])
    return _png(fig)


# ================================================================== PDF
def _estilos():
    ss = getSampleStyleSheet()
    return {
        "titulo": ParagraphStyle("t", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=24, textColor=colors.HexColor("#10203A"), alignment=TA_LEFT, spaceAfter=4),
        "sub": ParagraphStyle("s", parent=ss["Normal"], fontName="Helvetica", fontSize=11, textColor=colors.HexColor(CINZA), spaceAfter=14),
        "h": ParagraphStyle("h", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=13, textColor=colors.HexColor(AZUL_ESCURO), spaceBefore=8, spaceAfter=6),
        "p": ParagraphStyle("p", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13.5, textColor=colors.HexColor("#10203A"), spaceAfter=6),
        "peq": ParagraphStyle("q", parent=ss["Normal"], fontName="Helvetica", fontSize=7.5, textColor=colors.HexColor(CINZA)),
        "kpi_r": ParagraphStyle("kr", parent=ss["Normal"], fontName="Helvetica", fontSize=7.5, textColor=colors.HexColor(CINZA)),
        "kpi_v": ParagraphStyle("kv", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=15, textColor=colors.HexColor("#10203A"), leading=18),
        "cel": ParagraphStyle("c", parent=ss["Normal"], fontName="Helvetica", fontSize=8, leading=10),
    }


def _tabela(cab, linhas, larguras, e, alinh_dir=()):
    dados = [[Paragraph(f"<b>{c}</b>", e["cel"]) for c in cab]]
    for ln in linhas:
        dados.append([Paragraph(str(v), e["cel"]) for v in ln])
    t = Table(dados, colWidths=larguras, repeatRows=1)
    estilo = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF3FC")), ("LINEBELOW", (0, 0), (-1, 0), .6, colors.HexColor(AZUL)),
              ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFD")]),
              ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
    for c in alinh_dir:
        estilo.append(("ALIGN", (c, 1), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(estilo))
    return t


def _kpis(itens, e):
    rotulos = [Paragraph(r, e["kpi_r"]) for r, _ in itens]
    valores = [Paragraph(v, e["kpi_v"]) for _, v in itens]
    t = Table([rotulos, valores], colWidths=[26 * cm / max(len(itens), 1)] * len(itens))
    estilo = [("LEFTPADDING", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, 0), 7), ("BOTTOMPADDING", (0, 1), (-1, 1), 8),
              ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7FAFD"))]
    for i in range(len(itens)):
        estilo.append(("BOX", (i, 0), (i, 1), .6, colors.HexColor("#DCE5F0")))
    t.setStyle(TableStyle(estilo))
    return t


def _rodape(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor(CINZA))
    canvas.drawString(1.5 * cm, 1.0 * cm, f"{sobre.NOME} · {sobre.CREDITO_CURTO}")
    canvas.drawRightString(landscape(A4)[0] - 1.5 * cm, 1.0 * cm, f"Página {doc.page}")
    canvas.restoreState()


def gerar_pdf(dados: dict) -> bytes:
    """dados: dicionario com resumo, mensal, porServico, porCategoria, porNuvem, topRecursos, achados, previsao, otimizacao, alertas, filtros, moeda."""
    with _TRAVA_GRAFICOS:
        return _gerar_pdf(dados)


def _gerar_pdf(dados: dict) -> bytes:
    e = _estilos()
    m = dados.get("moeda", "USD")
    r = dados.get("resumo", {})
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.4 * cm, bottomMargin=1.6 * cm,
                            title=f"{sobre.NOME}: relatório executivo", author=sobre.AUTOR)
    el = []

    # capa
    gerado = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    el.append(Paragraph(f"{sobre.NOME}", e["titulo"]))
    el.append(Paragraph(f"Relatório executivo de custo em nuvem · gerado em {gerado} · {dados.get('filtrosTexto', 'todos os dados')}", e["sub"]))
    el.append(_kpis([("Custo total no período", _moeda(r.get("custoEfetivo", 0), m)),
                     ("Mês corrente", _moeda(r.get("mesAtual", 0), m)),
                     ("Projeção do mês", _moeda(r.get("projecaoMes", 0), m)),
                     ("Economia sobre a tabela", f"{_moeda(r.get('economia', 0), m)} ({r.get('percentualEconomia', 0):.1%})")], e))
    el.append(Spacer(1, 10))

    # narrativa
    el.append(Paragraph("Leitura executiva", e["h"]))
    for par in dados.get("narrativa", []):
        el.append(Paragraph(par, e["p"]))
    el.append(Spacer(1, 6))

    if dados.get("mensal"):
        el.append(Image(graf_mensal(dados["mensal"], m), width=22 * cm, height=6.6 * cm))
    el.append(PageBreak())

    # distribuicao
    el.append(Paragraph("Onde o dinheiro está", e["h"]))
    linha = []
    if dados.get("porServico"):
        linha.append(Image(graf_barras(dados["porServico"], "Maiores serviços"), width=14.5 * cm, height=5.2 * cm))
    if dados.get("porCategoria"):
        linha.append(Image(graf_rosca(dados["porCategoria"], "Por categoria"), width=11 * cm, height=5.2 * cm))
    if linha:
        el.append(Table([linha], colWidths=[15 * cm, 11.5 * cm]))
    if dados.get("porNuvem") and len(dados["porNuvem"]) > 1:
        el.append(Spacer(1, 6))
        el.append(Image(graf_barras(dados["porNuvem"], "Por nuvem", CORES_NUVEM, altura=2.2), width=25 * cm, height=5.6 * cm))
    el.append(Spacer(1, 6))
    if dados.get("topRecursos"):
        el.append(Paragraph("Recursos que mais consomem", e["h"]))
        el.append(_tabela(["Recurso", "Serviço", "Grupo", "Nuvem", "Custo"],
                          [[x["recurso"][:48], x["servico"][:30], x["grupo"][:26], x["nuvem"], _moeda(x["efetivo"], m)] for x in dados["topRecursos"][:12]],
                          [8.4 * cm, 5.8 * cm, 4.8 * cm, 3.6 * cm, 3.4 * cm], e, alinh_dir=(4,)))
    el.append(PageBreak())

    # previsao e otimizacao
    prev = dados.get("previsao") or {}
    if prev.get("previsao"):
        el.append(Paragraph("Previsão", e["h"]))
        el.append(_kpis([("Próximos 30 dias", _moeda(prev["acumulado"]["30"], m)), ("Próximos 60 dias", _moeda(prev["acumulado"]["60"], m)),
                         ("Próximos 90 dias", _moeda(prev["acumulado"]["90"], m)),
                         ("Tendência mensal", f"{prev.get('tendenciaMensal', 0):+.1%} · confiabilidade {prev.get('confiabilidade', '')}")], e))
        el.append(Spacer(1, 6))
        el.append(Image(graf_previsao(prev, m), width=25 * cm, height=8 * cm))
    ot = dados.get("otimizacao") or {}
    if ot.get("recomendacoes"):
        el.append(Paragraph(f"Oportunidades de economia: {_moeda(ot['resumo']['economiaEstimada'], m)} por período, em {len(ot['recomendacoes'])} recomendações", e["h"]))
        el.append(_tabela(["Categoria", "O que está em uso", "Sugestão", "Custo atual", "Economia estimada", "Confiança"],
                          [[x["categoria"], x["atual"][:44], x["sugerido"][:40], _moeda(x["custoAtual"], m), _moeda(x["economiaEstimada"], m), x["confianca"]]
                           for x in ot["recomendacoes"][:12]],
                          [3.2 * cm, 7.6 * cm, 6.6 * cm, 3 * cm, 3.2 * cm, 2.2 * cm], e, alinh_dir=(3, 4)))
    el.append(PageBreak())

    # insights e alertas
    if dados.get("achados"):
        el.append(Paragraph("Achados da análise automática", e["h"]))
        el.append(_tabela(["Severidade", "Achado", "Detalhe", "Valor"],
                          [[a["severidade"], a["titulo"], a["detalhe"][:170], _moeda(a["valor"], m) if a.get("valor") else ""] for a in dados["achados"][:12]],
                          [2.4 * cm, 6.4 * cm, 14 * cm, 3.2 * cm], e, alinh_dir=(3,)))
    if dados.get("alertas"):
        el.append(Spacer(1, 8))
        el.append(Paragraph("Alertas em aberto", e["h"]))
        el.append(_tabela(["Severidade", "Alerta", "Detalhe", "Aberto em"],
                          [[a["severidade"], a["titulo"], a["detalhe"][:150], (a.get("abertoEm") or "")[:10]] for a in dados["alertas"][:12]],
                          [2.4 * cm, 6.4 * cm, 14 * cm, 3.2 * cm], e))
    el.append(Spacer(1, 14))
    el.append(Paragraph(sobre.CREDITO_LONGO, e["peq"]))
    el.append(Paragraph(f"Método: agregações sobre dados FOCUS do FinOps hub; previsão por regressão linear com sazonalidade semanal; "
                        f"achados por regras estatísticas determinísticas, sem uso de modelo de linguagem. Estimativas de economia são referências de mercado.", e["peq"]))

    doc.build(el, onFirstPage=_rodape, onLaterPages=_rodape)
    return buf.getvalue()


# ================================================================== Excel
def gerar_excel(dados: dict) -> bytes:
    m = dados.get("moeda", "USD")
    wb = Workbook()
    cab_fill = PatternFill("solid", fgColor="0078D4")
    cab_font = Font(bold=True, color="FFFFFF", name="Segoe UI", size=10)
    borda = Border(bottom=Side(style="thin", color="DCE5F0"))
    fmt_moeda = '#,##0.00'
    fmt_pct = '0.0%'

    def aba(nome, cab, linhas, formatos=None):
        ws = wb.create_sheet(title=nome[:31])
        ws.append(cab)
        for c in ws[1]:
            c.fill, c.font, c.alignment = cab_fill, cab_font, Alignment(vertical="center")
        for ln in linhas:
            ws.append(ln)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border = borda
                c.font = Font(name="Segoe UI", size=10)
        for col, fmt in (formatos or {}).items():
            for c in ws[get_column_letter(col)][1:]:
                c.number_format = fmt
        for i, _ in enumerate(cab, start=1):
            largura = max(len(str(cab[i - 1])), *(len(str(ln[i - 1])) for ln in linhas[:300] if i - 1 < len(ln))) if linhas else len(str(cab[i - 1]))
            ws.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 10), 60)
        ws.freeze_panes = "A2"
        return ws

    r = dados.get("resumo", {})
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = sobre.NOME
    ws["A1"].font = Font(bold=True, size=16, color="0078D4", name="Segoe UI")
    ws["A2"] = f"Relatório gerado em {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M UTC')} · {dados.get('filtrosTexto', 'todos os dados')}"
    ws["A2"].font = Font(size=10, color="52658A", name="Segoe UI")
    linhas = [("Custo efetivo no período", r.get("custoEfetivo", 0)), ("Custo faturado", r.get("custoFaturado", 0)),
              ("Custo de tabela", r.get("custoLista", 0)), ("Economia", r.get("economia", 0)),
              ("Mês corrente", r.get("mesAtual", 0)), ("Projeção do mês", r.get("projecaoMes", 0)),
              ("Recursos com custo", r.get("recursos", 0)), ("Serviços", r.get("servicos", 0)), ("Nuvens", r.get("nuvens", 0)), ("Moeda", m)]
    for i, (k, v) in enumerate(linhas, start=4):
        ws.cell(row=i, column=1, value=k).font = Font(name="Segoe UI", size=10, bold=True)
        c = ws.cell(row=i, column=2, value=v)
        c.font = Font(name="Segoe UI", size=10)
        if isinstance(v, float):
            c.number_format = fmt_moeda
    ws.cell(row=len(linhas) + 6, column=1, value=sobre.CREDITO_LONGO).font = Font(name="Segoe UI", size=9, color="52658A", italic=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22

    aba("Mensal", ["Mês", "Efetivo", "Faturado", "Tabela", "Economia"],
        [[x["mes"], x["efetivo"], x["faturado"], x["lista"], x["economia"]] for x in dados.get("mensal", [])], {2: fmt_moeda, 3: fmt_moeda, 4: fmt_moeda, 5: fmt_moeda})
    aba("Diário", ["Data", "Efetivo", "Faturado"], [[x["data"], x["efetivo"], x["faturado"]] for x in dados.get("diario", [])], {2: fmt_moeda, 3: fmt_moeda})
    for nome, chave in (("Por serviço", "porServico"), ("Por categoria", "porCategoria"), ("Por nuvem", "porNuvem"), ("Por região", "porRegiao"), ("Por assinatura", "porAssinatura")):
        if dados.get(chave):
            aba(nome, ["Nome", "Efetivo", "Faturado", "Lançamentos"], [[x["nome"], x["efetivo"], x["faturado"], x["linhas"]] for x in dados[chave]], {2: fmt_moeda, 3: fmt_moeda})
    if dados.get("topRecursos"):
        aba("Recursos", ["Recurso", "Serviço", "Grupo", "Região", "Nuvem", "Ambiente", "Efetivo", "Economia"],
            [[x["recurso"], x["servico"], x["grupo"], x["regiao"], x["nuvem"], x.get("ambiente", ""), x["efetivo"], x["economia"]] for x in dados["topRecursos"]], {7: fmt_moeda, 8: fmt_moeda})
    cb = dados.get("chargeback") or {}
    if cb.get("centros"):
        aba("Chargeback", ["Centro de custo", "Responsável", "E-mail", "Custo direto", "Rateio", "A cobrar", "Participação", "Orçamento mensal", "Gasto no mês"],
            [[x["centro"], x["responsavel"], x["email"], x["direto"], x["rateio"], x["cobrar"], x["participacao"], x["orcamentoMensal"], x["gastoMes"]] for x in cb["centros"]],
            {4: fmt_moeda, 5: fmt_moeda, 6: fmt_moeda, 7: fmt_pct, 8: fmt_moeda, 9: fmt_moeda})
    gv = dados.get("governanca") or {}
    if gv.get("cobertura"):
        aba("Tags cobertura", ["Chave", "Custo coberto", "Cobertura", "Valores distintos"], [[x["chave"], x["efetivo"], x["cobertura"], x["valoresDistintos"]] for x in gv["cobertura"]], {2: fmt_moeda, 3: fmt_pct})
    if gv.get("semTag"):
        aba("Sem tag", ["Recurso", "Serviço", "Grupo", "Assinatura", "Nuvem", "Custo"], [[x["recurso"], x["servico"], x["grupo"], x["assinatura"], x["nuvem"], x["efetivo"]] for x in gv["semTag"]], {6: fmt_moeda})
    ot = dados.get("otimizacao") or {}
    if ot.get("recomendacoes"):
        aba("Otimização", ["Categoria", "Título", "Em uso", "Sugestão", "Recurso", "Serviço", "Nuvem", "Custo atual", "Economia mín", "Economia máx", "Economia estimada", "Confiança", "Ação"],
            [[x["categoria"], x["titulo"], x["atual"], x["sugerido"], x["recurso"], x["servico"], x["nuvem"], x["custoAtual"], x["economiaMin"], x["economiaMax"], x["economiaEstimada"], x["confianca"], x["acao"]] for x in ot["recomendacoes"]],
            {8: fmt_moeda, 9: fmt_moeda, 10: fmt_moeda, 11: fmt_moeda})
    prev = dados.get("previsao") or {}
    if prev.get("previsao"):
        aba("Previsão", ["Data", "Previsto", "Mínimo", "Máximo"], [[x["data"], x["valor"], x["min"], x["max"]] for x in prev["previsao"]], {2: fmt_moeda, 3: fmt_moeda, 4: fmt_moeda})
    if dados.get("previsaoWorkloads"):
        aba("Previsão por workload", ["Workload", "Últimos 30", "Próximos 30", "Próximos 60", "Próximos 90", "Tendência mensal", "Confiabilidade"],
            [[x["nome"], x["ultimos30"], x["proximos30"], x["proximos60"], x["proximos90"], x["tendenciaMensal"], x["confiabilidade"]] for x in dados["previsaoWorkloads"]],
            {2: fmt_moeda, 3: fmt_moeda, 4: fmt_moeda, 5: fmt_moeda, 6: fmt_pct})
    ia = dados.get("ia") or {}
    if ia.get("porServico"):
        aba("IA por serviço", ["Serviço", "Efetivo", "Faturado", "Lançamentos"], [[x["nome"], x["efetivo"], x["faturado"], x["linhas"]] for x in ia["porServico"]], {2: fmt_moeda, 3: fmt_moeda})
        aba("IA por modelo", ["Modelo", "Efetivo", "Faturado", "Lançamentos"], [[x["nome"], x["efetivo"], x["faturado"], x["linhas"]] for x in ia.get("porModelo", [])], {2: fmt_moeda, 3: fmt_moeda})
    bd = dados.get("bancos") or {}
    if bd.get("porEngine"):
        aba("Bancos por engine", ["Engine", "Efetivo", "Faturado", "Lançamentos"], [[x["nome"], x["efetivo"], x["faturado"], x["linhas"]] for x in bd["porEngine"]], {2: fmt_moeda, 3: fmt_moeda})
    if dados.get("achados"):
        aba("Insights", ["Severidade", "Categoria", "Título", "Detalhe", "Valor", "Ação"], [[a["severidade"], a["categoria"], a["titulo"], a["detalhe"], a.get("valor") or 0, a.get("acao") or ""] for a in dados["achados"]], {5: fmt_moeda})
    if dados.get("alertas"):
        aba("Alertas", ["Estado", "Severidade", "Regra", "Título", "Detalhe", "Valor", "Aberto em"], [[a["estado"], a["severidade"], a["regraNome"], a["titulo"], a["detalhe"], a.get("valor") or 0, a.get("abertoEm", "")] for a in dados["alertas"]], {6: fmt_moeda})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def narrativa(resumo: dict, achados: list, prev: dict, ot: dict, moeda: str) -> list[str]:
    """Paragrafos em portugues para a leitura executiva do PDF."""
    paras = []
    r = resumo
    if r.get("custoEfetivo"):
        p = (f"O custo efetivo do período foi de {_moeda(r['custoEfetivo'], moeda)}, distribuído em {r.get('servicos', 0)} serviços e "
             f"{r.get('recursos', 0)} recursos em {r.get('nuvens', 1)} nuvem(ns). ")
        if r.get("mesAnterior"):
            v = r.get("variacaoMensal", 0)
            comp = (f"Comparando os primeiros {r['diasComparados']} dias de cada mês, " if r.get("comparacaoParcial") else "")
            p += f"{comp}o mês corrente está {abs(v):.1%} {'acima' if v > 0 else 'abaixo'} do mês anterior, com projeção de fechamento em {_moeda(r.get('projecaoMes', 0), moeda)}. "
        if r.get("economia"):
            p += f"Os descontos já obtidos representam {r.get('percentualEconomia', 0):.1%} sobre o preço de tabela, ou {_moeda(r['economia'], moeda)}."
        paras.append(p)
    criticos = [a for a in achados if a["severidade"] == "critico"]
    atencao = [a for a in achados if a["severidade"] == "atencao"]
    if criticos or atencao:
        p = "A análise automática destacou "
        partes = []
        if criticos:
            partes.append(f"{len(criticos)} ponto(s) crítico(s), sendo o principal: {criticos[0]['titulo'].lower()}")
        if atencao:
            partes.append(f"{len(atencao)} ponto(s) de atenção, como {atencao[0]['titulo'].lower()}")
        paras.append(p + " e ".join(partes) + ".")
    if prev and prev.get("acumulado", {}).get("30"):
        paras.append(f"Mantido o comportamento atual, a previsão é de {_moeda(prev['acumulado']['30'], moeda)} nos próximos 30 dias e "
                     f"{_moeda(prev['acumulado']['90'], moeda)} em 90 dias, com tendência de {prev.get('tendenciaMensal', 0):+.1%} ao mês "
                     f"e confiabilidade {prev.get('confiabilidade', 'média')}.")
    if ot and ot.get("resumo", {}).get("economiaEstimada"):
        paras.append(f"Foram identificadas {ot['resumo']['recomendacoes']} oportunidades de otimização, com economia estimada entre "
                     f"{_moeda(ot['resumo']['economiaMin'], moeda)} e {_moeda(ot['resumo']['economiaMax'], moeda)} no período, "
                     f"o equivalente a cerca de {ot['resumo']['percentual']:.1%} do gasto. As de alta confiança somam {_moeda(ot['resumo']['altaConfianca'], moeda)}.")
    if not paras:
        paras.append("Não há dados suficientes no período selecionado para uma leitura executiva.")
    return paras
