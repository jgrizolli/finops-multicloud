"""
FinOps Multicloud
Construido por Wanderlei Grizolli Junior, Sr. Solution Engineer.
Baseado no Microsoft FinOps toolkit (https://github.com/microsoft/finops-toolkit).

Teste local de todos os modulos, com dado sintetico. Nao precisa de Azure.
Uso:  python api/test_local.py
"""

from __future__ import annotations

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import alerts  # noqa: E402
import analytics  # noqa: E402
import chargeback  # noqa: E402
import demo_data  # noqa: E402
import domains  # noqa: E402
import export_report  # noqa: E402
import forecast  # noqa: E402
import governance  # noqa: E402
import insights  # noqa: E402
import optimization  # noqa: E402
from data_source import normalizar  # noqa: E402
from state_store import LocalJsonStore  # noqa: E402

falhas: list[str] = []


def secao(t):
    print(f"\n{'=' * 78}\n  {t}\n{'=' * 78}")


def ok(cond, msg):
    (print(f"  OK   {msg}") if cond else falhas.append(msg))
    if not cond:
        print(f"  X    {msg}")


def main() -> int:
    df, tags = normalizar(demo_data.gerar())
    r = analytics.resumo(df)
    m = r["moeda"]
    print(f"Dado de teste: {len(df):,} linhas, {df['Mes'].nunique()} meses, {df['Nuvem'].nunique()} nuvens")

    secao("normalizacao")
    ok(set(df["Nuvem"].unique()) == {"Microsoft Azure", "Amazon Web Services", "Google Cloud", "Oracle Cloud"}, "quatro nuvens canonicas")
    ok("SkuMeter" in df.columns and "ResourceGroupName" in df.columns, "colunas equivalentes renomeadas")
    ok(df["Ambiente"].isin(["Produção", "Não produção", "Desconhecido"]).all(), "ambiente classificado")
    ok((df["Ambiente"] == "Não produção").sum() > 0, "existem linhas de nao producao")
    ok(0 < df["SemTag"].mean() < 1, f"sem tag em {df['SemTag'].mean():.0%} das linhas")

    secao("fonte Kusto (sem conexao): consulta e normalizacao de Tags como dict")
    import kusto_source
    consulta = kusto_source.montar_consulta("Costs_v1_2()", 13)
    ok("set notruncation" in consulta and "Costs_v1_2()" in consulta and "startofmonth(now(), -12)" in consulta, "KQL montado com notruncation e janela de 13 meses")
    ok("EffectiveCost" in consulta and "x_SourceProvider" in consulta and "Tags" in consulta, "KQL projeta as colunas FOCUS necessarias")
    import pandas as _pd
    mini = demo_data.gerar(dias=20).head(50).copy()
    mini["Tags"] = [{"CostCenter": "CC-9", "Environment": "prod"}] * len(mini)  # como o Kusto devolve (dynamic -> dict)
    mini_n, mini_tags = normalizar(mini)
    ok((~mini_n["SemTag"]).all() and any("CostCenter" in v for v in mini_tags.values()), "Tags em formato dict (Kusto) sao lidas igual ao JSON do parquet")
    ok((mini_n["Ambiente"] == "Produção").all(), "ambiente classificado pela tag Environment vinda do Kusto")

    secao("resumo e comparacao honesta")
    ok(r["custoEfetivo"] > 0, f"custo efetivo {r['custoEfetivo']:,.2f}")
    ok(abs(r["variacaoMensal"]) < 1.5, f"variacao mensal plausivel ({r['variacaoMensal']:+.1%}, parcial={r['comparacaoParcial']})")

    secao("governanca")
    g = governance.analisar(df, tags, None, None, demo_data.ORCAMENTOS_DEMO)
    ok(g["chave"] == "CostCenter", f"chave padrao detectada: {g['chave']}")
    ok(len(g["porValor"]) >= 4, f"{len(g['porValor'])} valores de {g['chave']}")
    ok(len(g["semTag"]) > 0, f"{len(g['semTag'])} recursos sem tag listados")
    ok(len(g["orcamentosTag"]) == 2, f"{len(g['orcamentosTag'])} orcamentos por tag avaliados")
    ok(len(g["conformidade"]) == 3, "conformidade das 3 tags obrigatorias")

    secao("IA")
    ia = domains.analisar_ia(df, sorted(df["Mes"].unique()))
    ok(not ia["vazio"], "dominio de IA identificado")
    ok(ia["resumo"]["tokens"] > 0, f"tokens estimados: {ia['resumo']['tokens']:,.0f}")
    ok(any(x["nome"].startswith("GPT-4o") for x in ia["porModelo"]), f"modelos: {[x['nome'] for x in ia['porModelo'][:5]]}")
    ok(ia["resumo"]["custoAgentes"] > 0, f"agentes: {ia['resumo']['custoAgentes']:,.2f}")
    ok(len(ia["porNuvem"]) >= 3, f"IA em {len(ia['porNuvem'])} nuvens")

    secao("bancos de dados")
    bd = domains.analisar_bancos(df, sorted(df["Mes"].unique()))
    ok(not bd["vazio"], "dominio de bancos identificado")
    engines = [x["nome"] for x in bd["porEngine"]]
    ok("Azure SQL Database" in engines and "Amazon RDS" in engines, f"engines: {engines[:6]}")
    ok(len(bd["porTipo"]) >= 3, f"tipos: {[x['nome'] for x in bd['porTipo']]}")

    secao("otimizacao")
    ot = optimization.analisar(df)
    cats = {x["categoria"] for x in ot["recomendacoes"]}
    ok(len(ot["recomendacoes"]) >= 6, f"{len(ot['recomendacoes'])} recomendacoes")
    ok("Troca de tecnologia" in cats, "regras de troca de tecnologia dispararam")
    ok("Compromissos" in cats, "regras de compromisso dispararam")
    ok("Agendamento" in cats, "regra de fim de semana em nao producao disparou")
    for x in ot["recomendacoes"][:6]:
        print(f"       {x['categoria']:20} {x['titulo'][:48]:48} {x['economiaEstimada']:>10,.0f}  {x['confianca']}")

    secao("previsao")
    p = forecast.prever(df)
    ok(len(p["previsao"]) == 90, "90 dias de previsao")
    ok(p["acumulado"]["30"] > 0 and p["acumulado"]["90"] > p["acumulado"]["30"], f"acumulados 30/60/90: {p['acumulado']['30']:,.0f} / {p['acumulado']['60']:,.0f} / {p['acumulado']['90']:,.0f}")
    ok(all(x["min"] <= x["valor"] <= x["max"] for x in p["previsao"]), "faixa de confianca consistente")
    ok(p["confiabilidade"] in ("alta", "média", "baixa"), f"confiabilidade {p['confiabilidade']}, tendencia {p['tendenciaMensal']:+.1%}/mes")
    wl = forecast.prever_workloads(df, "ServiceName", top=5)
    ok(len(wl) == 5, f"previsao por workload: {[w['nome'] for w in wl]}")

    secao("chargeback")
    cb = chargeback.analisar(df, demo_data.CENTROS_DEMO, tags)
    nomes = [c["centro"] for c in cb["centros"]]
    ok("Plataforma" in nomes and chargeback.NAO_ALOCADO in nomes, f"centros: {nomes}")
    ok(cb["resumo"]["percentualAlocado"] > 0.5, f"{cb['resumo']['percentualAlocado']:.0%} alocado")
    soma = sum(c["cobrar"] for c in cb["centros"] if c["centro"] != chargeback.NAO_ALOCADO)
    ok(abs(soma - cb["resumo"]["total"]) < 1, "rateio fecha com o total")
    auto = chargeback.analisar(df, [], tags)
    ok(auto["automatico"] and len(auto["centros"]) >= 4, f"showback automatico por {auto['chaveAutomatica']}")
    csv = "nome,responsavel,email,orcamentoMensal,tipo,chave,valor\nRH,Bia,bia@x.com,1000,tag,CostCenter,CC-9\nRH,Bia,bia@x.com,1000,grupo,,rg-rh\n"
    imp = chargeback.importar(csv, "csv")
    ok(len(imp) == 1 and len(imp[0]["regras"]) == 2, "importacao CSV agrupa regras por centro")

    secao("alertas com persistencia local")
    with tempfile.TemporaryDirectory() as tmp:
        store = LocalJsonStore(tmp)
        for o in demo_data.ORCAMENTOS_DEMO:
            store.salvar("orcamentos", o)
        store.salvar("regras_alerta", {"id": "orc", "nome": "Orçamento total", "tipo": "orcamento", "limiar": 1000, "escopoTipo": "total",
                                       "severidade": "critico", "destinatarios": ["fin@x.com"], "ativo": True})
        res = alerts.avaliar(df, store, chargeback.alocar(df, demo_data.CENTROS_DEMO, tags), enviar_email=False)
        ok(res["regras"] >= 6, f"{res['regras']} regras (padrao semeadas + orcamento)")
        ok(len(res["novos"]) >= 3, f"{len(res['novos'])} alertas abertos")
        tipos = {a["tipo"] for a in res["novos"]}
        ok("orcamento" in tipos, "orcamento estourado detectado")
        ok("pico_diario" in tipos, "pico diario detectado")
        for a in res["novos"][:6]:
            print(f"       [{a['severidade']:11}] {a['titulo'][:60]}")
        res2 = alerts.avaliar(df, store, None, enviar_email=False)
        ok(len(res2["novos"]) == 0, "segunda avaliacao nao duplica")
        a0 = res["novos"][0]
        alerts.mudar_estado(store, a0["id"], "reconhecido", "vou tratar")
        ok(store.obter("alertas", a0["id"])["estado"] == "reconhecido", "mudanca de estado persistida")
        rs = alerts.resumo(store)
        ok(rs["abertos"] == len(res["novos"]) - 1, f"resumo: {rs['abertos']} abertos, {rs['reconhecidos']} reconhecido")

    secao("insights")
    ach = insights.gerar(df, m)
    ok(len(ach) >= 5, f"{len(ach)} achados")
    ok(any(a["categoria"] == "Anomalia" for a in ach), "pico artificial detectado como anomalia")
    ok(any(a["categoria"] == "Multicloud" for a in ach), "achado multicloud")

    secao("exportacao")
    prev = forecast.prever(df)
    dados = {"moeda": m, "resumo": r, "mensal": analytics.serie_mensal(df), "diario": analytics.serie_diaria(df),
             "porServico": analytics.agrupar(df, "ServiceName", 15), "porCategoria": analytics.agrupar(df, "ServiceCategory", 10),
             "porNuvem": analytics.agrupar(df, "Nuvem", 6), "porRegiao": analytics.agrupar(df, "RegionName", 12),
             "porAssinatura": analytics.agrupar(df, "SubAccountName", 15), "topRecursos": analytics.top_recursos(df, 50),
             "achados": ach, "previsao": prev, "otimizacao": ot, "alertas": [], "filtrosTexto": "teste",
             "chargeback": cb, "governanca": g, "previsaoWorkloads": wl, "ia": ia, "bancos": bd}
    dados["narrativa"] = export_report.narrativa(r, ach, prev, ot, m)
    pdf = export_report.gerar_pdf(dados)
    ok(pdf[:4] == b"%PDF" and len(pdf) > 30000, f"PDF gerado: {len(pdf) / 1024:,.0f} KB")
    xlsx = export_report.gerar_excel(dados)
    ok(xlsx[:2] == b"PK" and len(xlsx) > 20000, f"Excel gerado: {len(xlsx) / 1024:,.0f} KB")
    from openpyxl import load_workbook
    import io as _io
    wb = load_workbook(_io.BytesIO(xlsx))
    ok(len(wb.sheetnames) >= 15, f"{len(wb.sheetnames)} abas: {wb.sheetnames[:8]}...")
    saida = os.path.join(tempfile.gettempdir(), "finops-teste.pdf")
    open(saida, "wb").write(pdf)
    print(f"       PDF de amostra em {saida}")

    print()
    if falhas:
        print(f"FALHAS ({len(falhas)}):")
        for f in falhas:
            print(f"  - {f}")
        return 1
    print("Todos os testes passaram.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
